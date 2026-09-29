from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from .experiment_io_v2 import (
    assert_blind_records,
    read_jsonl,
    sha256_path,
    sha256_json,
    utc_now,
    write_json,
    write_jsonl,
)
from .extractor_v2 import CANDIDATE_GENERATOR_VERSION, RuleCandidateExtractorV2
from .io_utils import normalize_space


ANALYSIS_SCHEMA_VERSION = "ecospec-analysis-v2.1"


def _entity_key(entity: dict[str, Any]) -> tuple[str, str]:
    return normalize_space(str(entity.get("name", ""))), str(
        entity.get("entity_type", "")
    )


def _relation_key(relation: dict[str, Any]) -> tuple[str, str, str, str, str]:
    return (
        normalize_space(str(relation.get("head_name", ""))),
        str(relation.get("head_type", "")),
        str(relation.get("relation_type", "")),
        normalize_space(str(relation.get("tail_name", ""))),
        str(relation.get("tail_type", "")),
    )


def _rows_by_unit(rows: list[dict[str, Any]], label: str) -> dict[str, dict[str, Any]]:
    output = {str(row["unit_id"]): row for row in rows}
    if len(output) != len(rows):
        raise ValueError(f"{label} unit ids are duplicated")
    return output


def _assert_allowed_gold(rows: list[dict[str, Any]], allowed: set[str]) -> None:
    splits = {str(row.get("split", "missing")) for row in rows}
    if not splits <= allowed:
        raise ValueError(
            "analysis accepts only explicitly allowed gold splits; "
            f"allowed={sorted(allowed)} found={sorted(splits)}"
        )


def _unit_text(unit: dict[str, Any]) -> str:
    values: list[str] = []
    for key in ("clause_text", "introduction", "interstitial_text", "adjacent_source_text"):
        value = unit.get(key)
        if value:
            values.append(str(value))
    for formula in unit.get("formulas", []):
        values.append(str(formula.get("expression_text", "")))
    for definition in unit.get("variable_definitions", []):
        values.append(
            " ".join(
                str(definition.get(key, ""))
                for key in ("symbol", "definition", "unit")
            )
        )
    cells = unit.get("cells")
    if isinstance(cells, dict):
        values.extend(str(value) for value in cells.values())
    elif isinstance(cells, list):
        values.extend(str(value) for value in cells)
    if not values:
        values.extend(
            str(span.get("text", ""))
            for span in unit.get("provenance", {}).get("evidence_spans", [])
        )
    return normalize_space(" ".join(value for value in values if value))


def _span_ids(item: dict[str, Any]) -> str:
    return ";".join(str(value) for value in item.get("evidence_span_ids", []))


def _detail_base(unit: dict[str, Any], error_type: str) -> dict[str, Any]:
    provenance = unit.get("provenance", {})
    return {
        "unit_id": str(unit["unit_id"]),
        "standard_code": str(provenance.get("standard_code", "")),
        "unit_type": str(unit.get("unit_type", "")),
        "section": str(provenance.get("section", "")),
        "pages": ";".join(str(value) for value in provenance.get("pages", [])),
        "error_type": error_type,
        "object_kind": "entity" if error_type.startswith("entity_") else "relation",
        "expected_side": "gold" if error_type.endswith("false_negative") else "prediction",
        "entity_name": "",
        "entity_type": "",
        "head_name": "",
        "head_type": "",
        "relation_type": "",
        "tail_name": "",
        "tail_type": "",
        "evidence_span_ids": "",
        "source_text": _unit_text(unit),
        "candidate_generator": CANDIDATE_GENERATOR_VERSION,
        "candidate_present": False,
        "candidate_id": "",
        "selection_result": "candidate_missing",
        "review_reason": "",
        "correction_action": "",
    }


def _counter_dict(values: Iterable[str]) -> dict[str, int]:
    return dict(sorted(Counter(values).items()))


def _load_analysis_inputs(
    units_path: Path,
    gold_path: Path,
    predictions_path: Path,
) -> tuple[
    dict[str, dict[str, Any]],
    dict[str, dict[str, Any]],
    dict[str, dict[str, Any]],
]:
    units = read_jsonl(units_path)
    gold = read_jsonl(gold_path)
    predictions = read_jsonl(predictions_path)
    assert_blind_records(units)
    _assert_allowed_gold(gold, {"train", "dev"})
    unit_by_id = _rows_by_unit(units, "source")
    gold_by_id = _rows_by_unit(gold, "gold")
    prediction_by_id = _rows_by_unit(predictions, "prediction")
    if set(unit_by_id) != set(gold_by_id) or set(unit_by_id) != set(prediction_by_id):
        raise ValueError("source, gold, and prediction unit ids are not identical")
    return unit_by_id, gold_by_id, prediction_by_id


def collect_error_details_v2(
    units_path: Path,
    gold_path: Path,
    predictions_path: Path,
) -> list[dict[str, Any]]:
    unit_by_id, gold_by_id, prediction_by_id = _load_analysis_inputs(
        units_path, gold_path, predictions_path
    )
    details: list[dict[str, Any]] = []
    for unit_id in sorted(unit_by_id):
        detail_start = len(details)
        unit = unit_by_id[unit_id]
        gold = gold_by_id[unit_id]
        prediction = prediction_by_id[unit_id]
        candidates = RuleCandidateExtractorV2().predict_unit(unit)
        actual_hash = sha256_json({k: candidates[k] for k in ("entities", "relations")})
        expected_hash = prediction.get("candidate_hash")
        if expected_hash and expected_hash != actual_hash:
            raise ValueError(
                f"candidate hash mismatch for {unit_id}: expected={expected_hash}, actual={actual_hash}; "
                "use the original candidate generator and source units; do not rewrite historical hashes"
            )
        verification = "hash_verified" if expected_hash else "missing_hash"
        gold_entities = {_entity_key(item): item for item in gold.get("entities", [])}
        pred_entities = {
            _entity_key(item): item for item in prediction.get("entities", [])
        }
        candidate_entities = {
            _entity_key(item): item for item in candidates.get("entities", [])
        }
        gold_relations = {
            _relation_key(item): item for item in gold.get("relations", [])
        }
        pred_relations = {
            _relation_key(item): item for item in prediction.get("relations", [])
        }
        candidate_relations = {
            _relation_key(item): item for item in candidates.get("relations", [])
        }
        for error_type, keys, source in (
            ("entity_false_negative", set(gold_entities) - set(pred_entities), gold_entities),
            ("entity_false_positive", set(pred_entities) - set(gold_entities), pred_entities),
        ):
            for key in sorted(keys):
                item = source[key]
                row = _detail_base(unit, error_type)
                row.update(
                    entity_name=key[0],
                    entity_type=key[1],
                    evidence_span_ids=_span_ids(item),
                )
                candidate = candidate_entities.get(key)
                row.update(
                    candidate_present=candidate is not None,
                    candidate_id=(str(candidate.get("entity_id", "")) if candidate else ""),
                    selection_result=(
                        "selected_false_positive"
                        if key in pred_entities
                        else "candidate_not_selected"
                        if candidate is not None
                        else "candidate_missing"
                    ),
                )
                details.append(row)
        for error_type, keys, source in (
            (
                "relation_false_negative",
                set(gold_relations) - set(pred_relations),
                gold_relations,
            ),
            (
                "relation_false_positive",
                set(pred_relations) - set(gold_relations),
                pred_relations,
            ),
        ):
            for key in sorted(keys):
                item = source[key]
                row = _detail_base(unit, error_type)
                row.update(
                    head_name=key[0],
                    head_type=key[1],
                    relation_type=key[2],
                    tail_name=key[3],
                    tail_type=key[4],
                    evidence_span_ids=_span_ids(item),
                )
                candidate = candidate_relations.get(key)
                row.update(
                    candidate_present=candidate is not None,
                    candidate_id=(str(candidate.get("relation_id", "")) if candidate else ""),
                    selection_result=(
                        "selected_false_positive"
                        if key in pred_relations
                        else "candidate_not_selected"
                        if candidate is not None
                        else "candidate_missing"
                    ),
                )
                details.append(row)
        for row in details[detail_start:]:
            row["candidate_verification"] = verification
            row["reconstructed_candidate_hash"] = actual_hash
            row["prediction_candidate_hash"] = expected_hash or ""
            if not expected_hash:
                row["candidate_present"] = None
                row["candidate_id"] = ""
                if row["selection_result"] != "selected_false_positive":
                    row["selection_result"] = "candidate_unverified"
    return details


def _style_workbook(workbook: Any) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5D73")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            longest = max(len(str(cell.value or "")) for cell in column[:200])
            sheet.column_dimensions[letter].width = min(max(longest + 2, 10), 55)


def _write_workbook(
    path: Path,
    sheets: list[tuple[str, list[dict[str, Any]]]],
) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError(
            'Excel output requires openpyxl; install with pip install -e ".[analysis]"'
        ) from exc
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        fields = list(rows[0]) if rows else ["message"]
        sheet.append(fields)
        if rows:
            for row in rows:
                sheet.append([row.get(field, "") for field in fields])
        else:
            sheet.append(["no records"])
    _style_workbook(workbook)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def analyze_errors_v2(
    units_path: Path,
    gold_path: Path,
    predictions_path: Path,
    out_dir: Path,
) -> dict[str, Any]:
    details = collect_error_details_v2(units_path, gold_path, predictions_path)
    relation_error_matrix: dict[str, dict[str, int]] = defaultdict(
        lambda: {"false_positive": 0, "false_negative": 0, "total": 0}
    )
    for row in details:
        if row["object_kind"] != "relation":
            continue
        counts = relation_error_matrix[row["relation_type"]]
        key = (
            "false_positive"
            if row["error_type"] == "relation_false_positive"
            else "false_negative"
        )
        counts[key] += 1
        counts["total"] += 1
    summary = {
        "schema_version": ANALYSIS_SCHEMA_VERSION,
        "analysis_type": "prediction_errors",
        "created_at": utc_now(),
        "inputs": {
            "units": str(units_path),
            "units_sha256": sha256_path(units_path),
            "gold": str(gold_path),
            "gold_sha256": sha256_path(gold_path),
            "predictions": str(predictions_path),
            "predictions_sha256": sha256_path(predictions_path),
        },
        "error_count": len(details),
        "by_error_type": _counter_dict(row["error_type"] for row in details),
        "candidate_verification_by_error_row": _counter_dict(row["candidate_verification"] for row in details),
        "entity_errors_by_type": _counter_dict(
            row["entity_type"] for row in details if row["object_kind"] == "entity"
        ),
        "relation_errors_by_type": _counter_dict(
            row["relation_type"] for row in details if row["object_kind"] == "relation"
        ),
        "relation_errors_by_type_and_direction": {
            key: value for key, value in sorted(relation_error_matrix.items())
        },
        "by_standard": _counter_dict(row["standard_code"] for row in details),
        "by_unit_type": _counter_dict(row["unit_type"] for row in details),
        "human_review_fields": ["review_reason", "correction_action"],
    }
    out_dir.mkdir(parents=True, exist_ok=True)
    write_json(out_dir / "error_summary.json", summary)
    write_jsonl(out_dir / "error_details.jsonl", details)
    summary_rows = [
        {"category": "error_type", "name": key, "count": value}
        for key, value in summary["by_error_type"].items()
    ] + [
        {"category": "relation_type", "name": key, "count": value}
        for key, value in summary["relation_errors_by_type"].items()
    ] + [
        {
            "category": "entity_type",
            "name": key,
            "count": value,
        }
        for key, value in summary["entity_errors_by_type"].items()
    ] + [
        {
            "category": "standard",
            "name": key,
            "count": value,
        }
        for key, value in summary["by_standard"].items()
    ] + [
        {
            "category": "unit_type",
            "name": key,
            "count": value,
        }
        for key, value in summary["by_unit_type"].items()
    ]
    relation_direction_rows = [
        {"relation_type": key, **value}
        for key, value in summary[
            "relation_errors_by_type_and_direction"
        ].items()
    ]
    _write_workbook(
        out_dir / "error_details.xlsx",
        [
            ("错误明细", details),
            ("汇总", summary_rows),
            ("关系FP_FN", relation_direction_rows),
        ],
    )
    return summary


def _counter_rows(
    gold: Counter[str],
    candidates: Counter[str],
    covered: Counter[str],
    negatives: Counter[str],
) -> list[dict[str, Any]]:
    rows = []
    for name in sorted(set(gold) | set(candidates)):
        gold_count = gold[name]
        covered_count = covered[name]
        rows.append(
            {
                "type": name,
                "gold_count": gold_count,
                "candidate_count": candidates[name],
                "covered_gold_count": covered_count,
                "missed_gold_count": gold_count - covered_count,
                "candidate_negative_count": negatives[name],
                "coverage": round(covered_count / gold_count, 6)
                if gold_count
                else None,
            }
        )
    return rows


def analyze_training_v2(
    units_path: Path,
    annotations_path: Path,
    out_dir: Path,
) -> dict[str, Any]:
    units = read_jsonl(units_path)
    annotations = read_jsonl(annotations_path)
    assert_blind_records(units)
    _assert_allowed_gold(annotations, {"train"})
    unit_by_id = _rows_by_unit(units, "source")
    annotation_by_id = _rows_by_unit(annotations, "annotation")
    if set(unit_by_id) != set(annotation_by_id):
        raise ValueError("training source and annotation unit ids are not identical")

    entity_gold: Counter[str] = Counter()
    entity_candidates: Counter[str] = Counter()
    entity_covered: Counter[str] = Counter()
    entity_negatives: Counter[str] = Counter()
    relation_gold: Counter[str] = Counter()
    relation_candidates: Counter[str] = Counter()
    relation_covered: Counter[str] = Counter()
    relation_negatives: Counter[str] = Counter()
    unit_types: Counter[str] = Counter()
    standards: Counter[str] = Counter()
    no_relation_units = 0
    duplicate_texts: defaultdict[str, list[str]] = defaultdict(list)
    extractor = RuleCandidateExtractorV2()

    for unit_id, unit in unit_by_id.items():
        annotation = annotation_by_id[unit_id]
        prediction = extractor.predict_unit(unit)
        unit_types[str(unit.get("unit_type", ""))] += 1
        standards[str(unit.get("provenance", {}).get("standard_code", ""))] += 1
        duplicate_texts[_unit_text(unit)].append(unit_id)
        if not annotation.get("relations"):
            no_relation_units += 1

        gold_entities = {_entity_key(item) for item in annotation.get("entities", [])}
        candidate_entities = {
            _entity_key(item) for item in prediction.get("entities", [])
        }
        for key in gold_entities:
            entity_gold[key[1]] += 1
            if key in candidate_entities:
                entity_covered[key[1]] += 1
        for key in candidate_entities:
            entity_candidates[key[1]] += 1
            if key not in gold_entities:
                entity_negatives[key[1]] += 1

        gold_relations = {
            _relation_key(item) for item in annotation.get("relations", [])
        }
        candidate_relations = {
            _relation_key(item) for item in prediction.get("relations", [])
        }
        for key in gold_relations:
            relation_gold[key[2]] += 1
            if key in candidate_relations:
                relation_covered[key[2]] += 1
        for key in candidate_relations:
            relation_candidates[key[2]] += 1
            if key not in gold_relations:
                relation_negatives[key[2]] += 1

    entity_rows = _counter_rows(
        entity_gold, entity_candidates, entity_covered, entity_negatives
    )
    relation_rows = _counter_rows(
        relation_gold, relation_candidates, relation_covered, relation_negatives
    )
    duplicate_groups = [ids for text, ids in duplicate_texts.items() if text and len(ids) > 1]
    summary = {
        "schema_version": ANALYSIS_SCHEMA_VERSION,
        "analysis_type": "training_distribution",
        "created_at": utc_now(),
        "candidate_generator": CANDIDATE_GENERATOR_VERSION,
        "inputs": {
            "units": str(units_path),
            "units_sha256": sha256_path(units_path),
            "annotations": str(annotations_path),
            "annotations_sha256": sha256_path(annotations_path),
        },
        "unit_count": len(units),
        "no_relation_unit_count": no_relation_units,
        "unit_type_counts": dict(sorted(unit_types.items())),
        "standard_counts": dict(sorted(standards.items())),
        "gold_entity_count": sum(entity_gold.values()),
        "covered_entity_count": sum(entity_covered.values()),
        "entity_candidate_negative_count": sum(entity_negatives.values()),
        "entity_recall_upper_bound": round(
            sum(entity_covered.values()) / sum(entity_gold.values()), 6
        )
        if entity_gold
        else 1.0,
        "gold_relation_count": sum(relation_gold.values()),
        "covered_relation_count": sum(relation_covered.values()),
        "relation_candidate_negative_count": sum(relation_negatives.values()),
        "relation_recall_upper_bound": round(
            sum(relation_covered.values()) / sum(relation_gold.values()), 6
        )
        if relation_gold
        else 1.0,
        "duplicate_source_text_group_count": len(duplicate_groups),
        "duplicate_source_text_record_count": sum(len(ids) for ids in duplicate_groups),
        "entities_by_type": entity_rows,
        "relations_by_type": relation_rows,
    }
    summary_rows = [
        {"metric": key, "value": value}
        for key, value in summary.items()
        if isinstance(value, (str, int, float))
    ]
    unit_rows = [
        {"category": "unit_type", "name": key, "count": value}
        for key, value in sorted(unit_types.items())
    ] + [
        {"category": "standard", "name": key, "count": value}
        for key, value in sorted(standards.items())
    ]
    duplicate_rows = [
        {"normalized_text": text, "unit_ids": ";".join(ids), "count": len(ids)}
        for text, ids in duplicate_texts.items()
        if text and len(ids) > 1
    ]
    out_dir.mkdir(parents=True, exist_ok=True)
    write_json(out_dir / "training_distribution.json", summary)
    _write_workbook(
        out_dir / "training_distribution.xlsx",
        [
            ("总览", summary_rows),
            ("实体类型", entity_rows),
            ("关系类型", relation_rows),
            ("单元分布", unit_rows),
            ("重复来源单元", duplicate_rows),
        ],
    )
    return summary


__all__ = [
    "ANALYSIS_SCHEMA_VERSION",
    "analyze_errors_v2",
    "analyze_training_v2",
    "collect_error_details_v2",
]
